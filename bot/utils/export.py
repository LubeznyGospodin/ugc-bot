"""Экспорт базы бота в ОДИН xlsx: лист на каждый срез (креаторы / заходы / не зарегались)."""
from __future__ import annotations

import io

from openpyxl import Workbook
from sqlalchemy import select

from bot.database import get_session
from bot.models import BotVisit, Creator

COLUMNS = [
    ("full_name", "Имя и фамилия"),
    ("telegram_contact", "Telegram"),
    ("instagram", "Instagram"),
    ("other_socials", "Другие соцсети"),
    ("rate", "Оплата"),
    ("portfolio", "Портфолио"),
    ("age", "Возраст"),
    ("city", "Город"),
    ("phone", "Телефон"),
    ("categories", "Категории"),
    ("created_at", "Добавлен"),
]


def _dt(value) -> str:
    return value.strftime("%Y-%m-%d %H:%M") if value else ""


async def export_all_xlsx() -> io.BytesIO:
    """Одна книга, три листа:
    «Креаторы» — все анкеты; «Заходы» — все, кто жал /start (с меткой источника);
    «Не зарегались» — заходили, но анкету не заполнили (ровно те, кому уйдёт пуш)."""
    async with get_session() as session:
        creators = (await session.execute(select(Creator).order_by(Creator.created_at))).scalars().all()
        visits = (await session.execute(select(BotVisit).order_by(BotVisit.first_seen))).scalars().all()
    registered = {c.tg_id for c in creators}

    wb = Workbook()
    ws = wb.active
    ws.title = "Креаторы"
    ws.append([label for _, label in COLUMNS])
    for c in creators:
        ws.append([str(getattr(c, field) or "") for field, _ in COLUMNS])

    ws = wb.create_sheet("Заходы")
    ws.append(["Telegram ID", "Username", "Имя", "Источник", "Первый заход", "Последний заход", "Зарегался"])
    for v in visits:
        ws.append([v.tg_id, str(v.username or ""), str(v.full_name or ""), str(v.source or ""),
                   _dt(v.first_seen), _dt(v.last_seen), "да" if v.tg_id in registered else "нет"])

    ws = wb.create_sheet("Не зарегались")
    ws.append(["Telegram ID", "Username", "Имя", "Первый заход", "Пуш отправлен"])
    for v in visits:
        if v.tg_id in registered:
            continue
        ws.append([v.tg_id, str(v.username or ""), str(v.full_name or ""),
                   _dt(v.first_seen), _dt(v.nudged_at) or "нет"])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
